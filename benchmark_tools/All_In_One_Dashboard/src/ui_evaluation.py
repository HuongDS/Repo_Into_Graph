import json
import re
import threading
from tkinter import filedialog, messagebox

import customtkinter as ctk
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import api
from . import theme as T

GUID_RE = re.compile(r"^[0-9a-fA-F]{8}-([0-9a-fA-F]{4}-){3}[0-9a-fA-F]{12}$")

# (giá trị gửi API, tên hiển thị, mô tả, nhóm trong thực nghiệm)
FORMATS = [
    ("Sandwich", "Sandwich", "Khen điểm đúng → chỉ lỗi → động viên", "A"),
    ("Rubric", "Rubric", "Chấm rành mạch theo từng tiêu chí", "B"),
    ("Explainable", "Explainable", "Trích dẫn code, giải thích vì sao sai", "C"),
]
FORMAT_VALUES = {f[0].lower() for f in FORMATS}
CONTEXT_MODES = ["Raw", "Graph", "Hybrid"]

KEY_LABELS = {
    "positives": "Điểm tốt",
    "constructive_criticism": "Điểm cần sửa",
    "next_steps": "Hướng cải thiện",
    "criteria": "Tiêu chí",
    "code_evidence": "Bằng chứng từ code",
    "logic_explanation": "Giải thích",
    "comment": "Nhận xét",
    "raw_response": "Phản hồi gốc (không phải JSON hợp lệ)",
}

TEMPLATE_COLS = ["BusinessId", "QuestionNo", "Question", "ReferenceMaterial",
                 "StudentAnswer", "MaxScore", "FeedbackFormat"]
RESULT_COLS = ["BusinessId", "QuestionNo", "Question", "ReferenceMaterial", "StudentAnswer",
               "FeedbackFormat", "ContextMode", "MaxScore", "Score", "Feedback", "RawJSON", "Error"]


def score_of(data):
    if not isinstance(data, dict):
        return None
    return data.get("score", data.get("total_score"))


def readable(data):
    if not isinstance(data, dict):
        return str(data)
    lines = []
    for k, v in data.items():
        if k in ("score", "total_score", "feedback_type"):
            continue
        lines.append(KEY_LABELS.get(k, k.replace("_", " ").capitalize()).upper())
        if isinstance(v, list):
            for item in v:
                if isinstance(item, dict):
                    name = item.get("name", "Tiêu chí")
                    sc = item.get("score")
                    lines.append(f"• {name}" + (f"  —  {sc} điểm" if sc is not None else ""))
                    if item.get("comment"):
                        lines.append(f"   {item['comment']}")
                else:
                    lines.append(f"• {item}")
        else:
            lines.append(str(v))
        lines.append("")
    return "\n".join(lines).strip()


def fmt_num(x):
    try:
        return f"{float(x):g}"
    except (TypeError, ValueError):
        return str(x)


class PageStudentEvaluation(ctk.CTkFrame):
    def __init__(self, parent, controller):
        super().__init__(parent, fg_color="transparent", corner_radius=0)
        self.controller = controller
        self.q_labels = []
        self.current_q_index = None
        self.last_data = None
        self.last_results = []
        self.busy = False
        self._build()
        self.refresh_question_picker()

    def ui(self, fn, *args, **kwargs):
        self.after(0, lambda: fn(*args, **kwargs))

    # ------------------------------------------------------------------ layout
    def _build(self):
        self.grid_columnconfigure(0, weight=0, minsize=400)
        self.grid_columnconfigure(1, weight=3)
        self.grid_columnconfigure(2, weight=4)
        self.grid_rowconfigure(1, weight=1)

        T.page_header(self, "Chấm điểm",
                      "Chọn câu hỏi do AI sinh, nhập câu trả lời của sinh viên rồi để AI chấm và nhận xét."
                      ).grid(row=0, column=0, columnspan=3, sticky="ew",
                             padx=T.SP["xl"], pady=(T.SP["lg"], T.SP["base"]))

        left = T.Card(self)
        left.grid(row=1, column=0, sticky="nsew", padx=(T.SP["xl"], T.SP["sm"]), pady=(0, T.SP["xl"]))
        mid = T.Card(self)
        mid.grid(row=1, column=1, sticky="nsew", padx=T.SP["sm"], pady=(0, T.SP["xl"]))
        right = T.Card(self)
        right.grid(row=1, column=2, sticky="nsew", padx=(T.SP["sm"], T.SP["xl"]), pady=(0, T.SP["xl"]))

        self._build_left(left)
        self._build_mid(mid)
        self._build_right(right)

    def _build_left(self, card):
        scroll = ctk.CTkScrollableFrame(card, fg_color="transparent", scrollbar_button_color=T.C["outline"])
        scroll.pack(fill="both", expand=True, padx=T.SP["xs"], pady=T.SP["xs"])
        body = ctk.CTkFrame(scroll, fg_color="transparent")
        body.pack(fill="both", expand=True, padx=T.SP["base"], pady=T.SP["base"])

        T.section_title(body, "Cấu hình chấm").pack(fill="x", pady=(0, T.SP["base"]))

        T.field_label(body, "Mã nghiệp vụ", "Business ID").pack(fill="x")
        self.ent_biz = T.entry(body, mono=True, placeholder_text="xxxxxxxx-xxxx-xxxx-xxxx-xxxxxxxxxxxx")
        self.ent_biz.pack(fill="x", pady=(T.SP["xs"] + 2, 0))
        self.lbl_biz_name = T.muted(body, "")
        self.lbl_biz_name.pack(fill="x", pady=(T.SP["xs"], 0))

        T.field_label(body, "Thang điểm tối đa").pack(fill="x", pady=(T.SP["md"], 0))
        self.ent_max = T.entry(body, width=110)
        self.ent_max.insert(0, "10")
        self.ent_max.pack(anchor="w", pady=(T.SP["xs"] + 2, 0))

        self.cfg_banner = T.Banner(body, wraplength=270)
        self.cfg_banner.attach(after=self.ent_max, fill="x", pady=(T.SP["sm"], 0))

        T.field_label(body, "Ngữ cảnh code gửi cho AI").pack(fill="x", pady=(T.SP["base"], 0))
        self.seg_mode = T.segmented(body, CONTEXT_MODES)
        self.seg_mode.set("Hybrid")
        self.seg_mode.pack(fill="x", pady=(T.SP["xs"] + 2, 0))

        T.field_label(body, "Kiểu nhận xét").pack(fill="x", pady=(T.SP["base"], T.SP["xs"]))
        self.format_var = ctk.StringVar(value="Sandwich")
        for value, name, desc, group in FORMATS:
            T.radio(body, f"{name}  ({group})", self.format_var, value).pack(anchor="w", pady=(T.SP["sm"], 0))
            T.muted(body, desc, wraplength=250).pack(fill="x", padx=(28, 0))

        self.btn_submit = T.primary_button(body, "Chấm điểm", self.on_submit)
        self.btn_submit.pack(fill="x", pady=(T.SP["lg"], 0))

        T.divider(body).pack(fill="x", pady=T.SP["lg"])

        T.section_title(body, "Chấm hàng loạt (Excel)").pack(fill="x")
        T.muted(body, "Template có sẵn các câu hỏi AI vừa sinh. Điền cột StudentAnswer; "
                      "cột FeedbackFormat để trống sẽ dùng kiểu đang chọn ở trên.",
                wraplength=270).pack(fill="x", pady=(T.SP["xs"], T.SP["md"]))
        self.btn_template = T.secondary_button(body, "Tải template kèm câu hỏi", self.export_template)
        self.btn_template.pack(fill="x")
        self.btn_import = T.secondary_button(body, "Import Excel & chấm", self.batch_evaluate)
        self.btn_import.pack(fill="x", pady=(T.SP["sm"], 0))
        self.btn_export = T.secondary_button(body, "Lưu kết quả ra Excel", self.export_result)
        self.btn_export.pack(fill="x", pady=(T.SP["sm"], 0))
        self.progress = ctk.CTkProgressBar(body, height=6, progress_color=T.C["primary"],
                                           fg_color=T.C["surface_alt"])
        self.progress.set(0)
        self.lbl_progress = T.muted(body, "")

    def _build_mid(self, card):
        card.grid_columnconfigure(0, weight=1)
        for r, w in ((4, 2), (6, 3), (8, 4)):
            card.grid_rowconfigure(r, weight=w)
        pad = T.SP["lg"]

        T.field_label(card, "Câu hỏi do AI sinh").grid(row=0, column=0, sticky="ew", padx=pad, pady=(pad, 0))
        pick = ctk.CTkFrame(card, fg_color="transparent")
        pick.grid(row=1, column=0, sticky="ew", padx=pad, pady=(T.SP["xs"] + 2, 0))
        pick.grid_columnconfigure(0, weight=1)
        self.cb_question = T.combobox(pick, [""], command=self.on_pick_question)
        self.cb_question.grid(row=0, column=0, sticky="ew")
        self.btn_goto_gen = T.text_button(pick, "Tạo câu hỏi", lambda: self.controller.show("gen"))
        self.btn_goto_gen.grid(row=0, column=1, padx=(T.SP["sm"], 0))
        self.lbl_q_meta = T.muted(card, "", wraplength=420)
        self.lbl_q_meta.grid(row=2, column=0, sticky="ew", padx=pad, pady=(T.SP["xs"], T.SP["sm"]))

        gap = T.SP["xs"] + 2
        T.field_label(card, "Câu hỏi").grid(row=3, column=0, sticky="ew", padx=pad, pady=(T.SP["sm"], gap))
        self.txt_question = T.textbox(card, height=70)
        self.txt_question.grid(row=4, column=0, sticky="nsew", padx=pad)

        ref_head = ctk.CTkFrame(card, fg_color="transparent")
        ref_head.grid(row=5, column=0, sticky="ew", padx=pad, pady=(T.SP["base"], gap))
        T.field_label(ref_head, "Đáp án mẫu").pack(side="left")
        T.chip(ref_head, "🔒 AI sinh · chỉ xem", "neutral").pack(side="left", padx=(T.SP["sm"], 0))
        # Đáp án mẫu chỉ lấy từ AI, người dùng không sửa được
        self.txt_ref = T.textbox(card, height=90, fg_color=T.C["bg"], border_color=T.C["outline"])
        self.txt_ref.grid(row=6, column=0, sticky="nsew", padx=pad)
        self._set_ref("")

        T.field_label(card, "Câu trả lời của sinh viên").grid(
            row=7, column=0, sticky="ew", padx=pad, pady=(T.SP["base"], gap))
        self.txt_answer = T.textbox(card, height=110, border_color=T.C["primary"], border_width=2)
        self.txt_answer.grid(row=8, column=0, sticky="nsew", padx=pad)

        self.ans_banner = T.Banner(card, wraplength=440)
        self.ans_banner.attach_grid(row=9, column=0, sticky="ew", padx=pad, pady=(T.SP["sm"], 0))
        ctk.CTkFrame(card, height=pad, width=1, fg_color="transparent").grid(row=10, column=0)

    def _build_right(self, card):
        card.grid_columnconfigure(0, weight=1)
        card.grid_rowconfigure(2, weight=1)
        pad = T.SP["lg"]

        head = ctk.CTkFrame(card, fg_color="transparent")
        head.grid(row=0, column=0, sticky="ew", padx=pad, pady=(pad - 4, 0))
        head.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(head, text="Kết quả", font=T.font("title"), text_color=T.C["on_surface"],
                     anchor="w").grid(row=0, column=0, sticky="w")
        self.seg_view = T.segmented(head, ["Dễ đọc", "JSON"], command=self._render_result_text, width=150)
        self.seg_view.set("Dễ đọc")
        self.seg_view.grid(row=0, column=1)

        self.score_strip = ctk.CTkFrame(card, fg_color=T.C["surface_alt"], corner_radius=T.R["md"])
        self.score_strip.grid(row=1, column=0, sticky="ew", padx=pad, pady=(T.SP["md"], 0))
        self.lbl_score = ctk.CTkLabel(self.score_strip, text="", font=T.font("display"),
                                      text_color=T.C["primary"])
        self.lbl_score.pack(side="left", padx=(T.SP["base"], T.SP["md"]), pady=T.SP["md"])
        info = ctk.CTkFrame(self.score_strip, fg_color="transparent")
        info.pack(side="left", fill="x", expand=True)
        self.chip_type = T.chip(info, "", "primary")
        self.chip_type.pack(anchor="w")
        self.lbl_result_meta = T.muted(info, "")
        self.lbl_result_meta.pack(anchor="w", pady=(T.SP["xs"], 0))
        self.score_strip.grid_remove()

        area = ctk.CTkFrame(card, fg_color="transparent")
        area.grid(row=2, column=0, sticky="nsew", padx=pad, pady=(T.SP["md"], pad))
        area.grid_columnconfigure(0, weight=1)
        area.grid_rowconfigure(0, weight=1)
        self.txt_result = T.textbox(area)
        self.state_view = T.StateView(area)
        for w in (self.txt_result, self.state_view):
            w.grid(row=0, column=0, sticky="nsew")
            w.grid_remove()
        self._show_state("🎓", "Chưa chấm câu nào",
                         "Chọn câu hỏi, nhập câu trả lời của sinh viên rồi bấm “Chấm điểm”.")

    # ------------------------------------------------------------- view state
    def _show_state(self, *args, **kw):
        self.txt_result.grid_remove()
        self.state_view.show(*args, **kw)
        self.state_view.grid()

    def _show_text(self, text):
        self.state_view.grid_remove()
        self.txt_result.configure(state="normal")
        self.txt_result.delete("1.0", "end")
        self.txt_result.insert("1.0", text)
        self.txt_result.grid()

    def _log(self, line):
        self.state_view.grid_remove()
        self.txt_result.grid()
        self.txt_result.insert("end", line + "\n")
        self.txt_result.see("end")

    # ------------------------------------------------------- question picker
    def on_show(self):
        pass

    @staticmethod
    def _q_label(i, q):
        text = " ".join(str(q.get("question", "")).split())
        if len(text) > 80:
            text = text[:80].rstrip() + "…"
        diff = q.get("difficulty") or "—"
        return f"Câu {i + 1} · {diff} — {text}"

    def refresh_question_picker(self, select=None):
        qs = self.controller.question_set.get("questions", [])
        if not qs:
            self.q_labels = []
            placeholder = "Chưa có câu hỏi — hãy sinh ở trang Tạo câu hỏi"
            self.cb_question.configure(values=[placeholder], state="normal")
            self.cb_question.set(placeholder)
            self.cb_question.configure(state="disabled")
            self.btn_goto_gen.grid()
            self.lbl_q_meta.configure(text="Bạn vẫn có thể nhập tay câu hỏi và Business ID bên dưới.")
            return
        self.q_labels = [self._q_label(i, q) for i, q in enumerate(qs)]
        self.cb_question.configure(values=self.q_labels, state="readonly")
        self.btn_goto_gen.grid_remove()
        qset = self.controller.question_set
        self.lbl_q_meta.configure(
            text=f"{qset.get('business_name') or 'Nghiệp vụ'}  ·  {qset.get('method') or ''}  ·  {len(qs)} câu")
        idx = select if select is not None and 0 <= select < len(qs) else 0
        self.cb_question.set(self.q_labels[idx])
        self._apply_question(idx)

    def _set_ref(self, text):
        self.txt_ref.configure(state="normal")
        self.txt_ref.delete("1.0", "end")
        self.txt_ref.insert("1.0", text or "(Chưa có đáp án mẫu — hãy chọn một câu hỏi do AI sinh.)")
        self.txt_ref.configure(state="disabled")

    def _ai_reference(self):
        """Đáp án mẫu gốc của câu đang chọn (không đọc từ ô hiển thị)."""
        qs = self.controller.question_set.get("questions", [])
        if self.current_q_index is not None and self.current_q_index < len(qs):
            return qs[self.current_q_index].get("suggestedAnswer", "")
        return ""

    def _ai_reference_map(self):
        qs = self.controller.question_set.get("questions", [])
        return {" ".join(str(q.get("question", "")).split()): q.get("suggestedAnswer", "") for q in qs}

    def on_pick_question(self, label):
        if label in self.q_labels:
            self._apply_question(self.q_labels.index(label))

    def _apply_question(self, idx):
        qset = self.controller.question_set
        q = qset["questions"][idx]
        self.current_q_index = idx
        self.ent_biz.delete(0, "end")
        self.ent_biz.insert(0, qset.get("business_id", ""))
        self.lbl_biz_name.configure(text=qset.get("business_name", ""))
        self.txt_question.delete("1.0", "end")
        self.txt_question.insert("1.0", q.get("question", ""))
        self._set_ref(q.get("suggestedAnswer", ""))
        self.ans_banner.hide()
        self.cfg_banner.hide()

    # ------------------------------------------------------------- single
    def _read_common(self):
        """Đọc + kiểm tra cấu hình chung. Trả về (biz_id, max_score) hoặc None nếu lỗi."""
        biz = self.ent_biz.get().strip()
        if not GUID_RE.match(biz):
            self.cfg_banner.show("error", "Business ID chưa đúng định dạng GUID. "
                                          "Hãy chọn câu hỏi do AI sinh để tự điền.")
            return None
        try:
            mx = float(self.ent_max.get().strip().replace(",", "."))
            if mx <= 0:
                raise ValueError
        except ValueError:
            self.cfg_banner.show("error", "Thang điểm phải là số lớn hơn 0.")
            return None
        self.cfg_banner.hide()
        return biz, mx

    def on_submit(self):
        if self.busy:
            return
        common = self._read_common()
        question = self.txt_question.get("1.0", "end-1c").strip()
        answer = self.txt_answer.get("1.0", "end-1c").strip()
        if not question:
            self.ans_banner.show("error", "Chưa có câu hỏi. Chọn câu hỏi do AI sinh hoặc nhập tay.")
            return
        if not answer:
            self.ans_banner.show("error", "Hãy nhập câu trả lời của sinh viên trước khi chấm.")
            self.txt_answer.focus_set()
            return
        self.ans_banner.hide()
        if not common:
            return
        biz, mx = common
        payload = {
            "businessId": biz,
            "question": question,
            "referenceMaterial": self._ai_reference(),
            "studentAnswer": answer,
            "contextModel": self.seg_mode.get().lower(),
            "feedbackFormat": self.format_var.get().lower(),
            "maxScore": mx,
        }
        q_no = (self.current_q_index + 1) if self.current_q_index is not None else ""
        self._set_busy(True, "Đang chấm…")
        self.score_strip.grid_remove()
        self._show_state("⏳", "AI đang chấm bài…", "Thường mất 5–30 giây.", loading=True)

        def work():
            try:
                data = api.evaluate(payload)
                self.ui(self._on_result, payload, data, q_no)
            except api.ApiError as e:
                self.ui(self._show_state, "⚠", "Chấm điểm không thành công", str(e), "Thử lại", self.on_submit)
            finally:
                self.ui(self._set_busy, False)

        threading.Thread(target=work, daemon=True).start()

    def _set_busy(self, busy, text=None):
        self.busy = busy
        state = "disabled" if busy else "normal"
        self.btn_submit.configure(state=state, text=text if busy and text else "Chấm điểm")
        for b in (self.btn_import, self.btn_template, self.btn_export):
            b.configure(state=state)

    def _on_result(self, payload, data, q_no):
        self.last_data = data
        sc = score_of(data)
        self.lbl_score.configure(text=f"{fmt_num(sc) if sc is not None else '—'} / {fmt_num(payload['maxScore'])}")
        ftype = data.get("feedback_type") if isinstance(data, dict) else None
        self.chip_type.configure(text=ftype or payload["feedbackFormat"].capitalize())
        self.lbl_result_meta.configure(
            text=f"{'Câu ' + str(q_no) + '  ·  ' if q_no else ''}Ngữ cảnh {payload['contextModel'].capitalize()}")
        self.score_strip.grid()
        self._render_result_text()
        self.last_results.append(self._result_row(payload, data, q_no))

    def _render_result_text(self, _value=None):
        if self.last_data is None:
            return
        if self.seg_view.get() == "JSON":
            self._show_text(json.dumps(self.last_data, indent=2, ensure_ascii=False))
        else:
            self._show_text(readable(self.last_data))

    @staticmethod
    def _result_row(payload, data, q_no, error=""):
        sc = score_of(data) if data is not None else None
        return {
            "BusinessId": payload.get("businessId", ""),
            "QuestionNo": q_no,
            "Question": payload.get("question", ""),
            "ReferenceMaterial": payload.get("referenceMaterial", ""),
            "StudentAnswer": payload.get("studentAnswer", ""),
            "FeedbackFormat": payload.get("feedbackFormat", ""),
            "ContextMode": payload.get("contextModel", ""),
            "MaxScore": payload.get("maxScore", ""),
            "Score": sc if sc is not None else ("ERROR" if error else ""),
            "Feedback": readable(data) if data is not None else "",
            "RawJSON": json.dumps(data, ensure_ascii=False) if data is not None else "",
            "Error": error,
        }

    # -------------------------------------------------------------- excel
    @staticmethod
    def _style_sheet(ws, widths):
        head_fill = PatternFill("solid", fgColor="E3E8FF")
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = head_fill
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"

    def export_template(self):
        qset = self.controller.question_set
        mx = self.ent_max.get().strip() or "10"
        wb = Workbook()
        ws = wb.active
        ws.title = "Answers"
        ws.append(TEMPLATE_COLS)
        if qset.get("questions"):
            for i, q in enumerate(qset["questions"], 1):
                ws.append([qset.get("business_id", ""), i, q.get("question", ""),
                           q.get("suggestedAnswer", ""), "", mx, ""])
            default_name = "Evaluation_Template_" + re.sub(r"[^\w]+", "_", qset.get("business_name") or "questions") + ".xlsx"
        else:
            ws.append(["00000000-0000-0000-0000-000000000000", 1, "Ví dụ câu hỏi?", "Ví dụ đáp án", "Em trả lời là…", mx, ""])
            default_name = "Evaluation_Template.xlsx"
        self._style_sheet(ws, {"A": 38, "B": 11, "C": 60, "D": 60, "E": 60, "F": 10, "G": 16})

        guide = wb.create_sheet("HuongDan")
        for line in [
            "Mỗi dòng = 1 câu trả lời của 1 sinh viên cho 1 câu hỏi.",
            "Có nhiều sinh viên? Sao chép dòng câu hỏi rồi điền StudentAnswer khác nhau.",
            "Dòng có StudentAnswer trống sẽ được bỏ qua.",
            "ReferenceMaterial chỉ để xem: khi chấm, app lấy đáp án mẫu gốc do AI sinh (sửa cột này không có tác dụng).",
            "FeedbackFormat: Sandwich (A) / Rubric (B) / Explainable (C). Để trống = dùng kiểu đang chọn trong app.",
            "Có thể thêm cột riêng (vd. StudentId, Group) — các cột này được giữ nguyên và chép sang file kết quả.",
        ]:
            guide.append([line])
        guide.column_dimensions["A"].width = 110

        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
                                            initialfile=default_name)
        if not path:
            return
        try:
            wb.save(path)
            self.cfg_banner.hide()
            messagebox.showinfo("Đã lưu template", f"Đã lưu tại:\n{path}")
        except OSError as e:
            messagebox.showerror("Không lưu được", f"{e}\n\nNếu file đang mở trong Excel, hãy đóng lại rồi thử lại.")

    def batch_evaluate(self):
        if self.busy:
            return
        path = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])
        if not path:
            return
        try:
            ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
            rows = list(ws.iter_rows(values_only=True))
        except Exception as e:
            messagebox.showerror("Không đọc được file", str(e))
            return
        if len(rows) < 2:
            messagebox.showwarning("File trống", "File Excel không có dòng dữ liệu nào.")
            return
        idx = {str(h).strip().lower(): i for i, h in enumerate(rows[0]) if h is not None}
        missing = [c for c in ("BusinessId", "Question", "StudentAnswer") if c.lower() not in idx]
        if missing:
            messagebox.showerror("Thiếu cột", "File thiếu cột: " + ", ".join(missing) +
                                 "\nHãy dùng nút “Tải template kèm câu hỏi”.")
            return

        try:
            default_max = float(self.ent_max.get().strip().replace(",", ".") or 10)
        except ValueError:
            default_max = 10.0
        ui_format = self.format_var.get().lower()
        mode = self.seg_mode.get().lower()

        known = {c.lower() for c in TEMPLATE_COLS}
        extra_cols = [(str(h).strip(), i) for i, h in enumerate(rows[0])
                      if h is not None and str(h).strip().lower() not in known]
        ai_refs = self._ai_reference_map()
        jobs, skipped, ref_warn = [], 0, []
        for n, r in enumerate(rows[1:], start=2):
            def get(col):
                i = idx.get(col.lower())
                return "" if i is None or i >= len(r) or r[i] is None else str(r[i]).strip()
            if not get("StudentAnswer"):
                skipped += 1
                continue
            try:
                mx = float(get("MaxScore").replace(",", ".")) if get("MaxScore") else default_max
            except ValueError:
                mx = default_max
            fmt = get("FeedbackFormat").lower() or ui_format
            extras = {name: (r[i] if i < len(r) else None) for name, i in extra_cols}
            # Đáp án mẫu luôn lấy từ bộ câu hỏi AI sinh; cột trong Excel bị bỏ qua nếu khớp được câu hỏi
            key = " ".join(get("Question").split())
            if key in ai_refs:
                ai_ref = ai_refs[key]
            else:
                ai_ref = get("ReferenceMaterial")
                ref_warn.append(n)
            jobs.append((n, get("QuestionNo"), fmt, extras, {
                "businessId": get("BusinessId"), "question": get("Question"),
                "referenceMaterial": ai_ref, "studentAnswer": get("StudentAnswer"),
                "contextModel": mode, "feedbackFormat": fmt, "maxScore": mx,
            }))
        if not jobs:
            messagebox.showwarning("Không có gì để chấm", "Tất cả dòng đều trống cột StudentAnswer.")
            return

        self.last_results = []
        self.last_data = None
        self.score_strip.grid_remove()
        self._show_text("")
        self._log(f"Bắt đầu chấm {len(jobs)} câu trả lời (bỏ qua {skipped} dòng trống) — ngữ cảnh {mode}.")
        if ref_warn:
            self._log(f"⚠ {len(ref_warn)} dòng có câu hỏi không nằm trong bộ câu hỏi AI vừa sinh "
                      f"(dòng {', '.join(map(str, ref_warn[:10]))}{'…' if len(ref_warn) > 10 else ''}) — "
                      "dùng đáp án mẫu ghi trong file. Hãy sinh lại đúng bộ câu hỏi trước khi import để chắc chắn.")
        self._set_busy(True, "Đang chấm hàng loạt…")
        self.progress.set(0)
        self.progress.pack(fill="x", pady=(T.SP["md"], 0))
        self.lbl_progress.pack(fill="x", pady=(T.SP["xs"], 0))
        self.lbl_progress.configure(text=f"0 / {len(jobs)}")

        def work():
            ok = fail = 0
            for k, (row_no, q_no, fmt, extras, payload) in enumerate(jobs, 1):
                err, data = "", None
                if not GUID_RE.match(payload["businessId"]):
                    err = "BusinessId không hợp lệ"
                elif fmt not in FORMAT_VALUES:
                    err = f"FeedbackFormat '{fmt}' không hợp lệ"
                else:
                    try:
                        data = api.evaluate(payload)
                    except api.ApiError as e:
                        err = str(e)
                if err:
                    fail += 1
                    self.ui(self._log, f"✗ Dòng {row_no}: {err}")
                else:
                    ok += 1
                    sc = score_of(data)
                    self.ui(self._log, f"✓ Dòng {row_no} · {fmt}: {fmt_num(sc) if sc is not None else '?'} / {fmt_num(payload['maxScore'])}")
                row = dict(extras)
                row.update(self._result_row(payload, data, q_no, err))
                self.last_results.append(row)
                self.ui(self.progress.set, k / len(jobs))
                self.ui(self.lbl_progress.configure, text=f"{k} / {len(jobs)}")
            self.ui(self._log, f"\nXong: {ok} thành công, {fail} lỗi. Bấm “Lưu kết quả ra Excel” để xuất file.")
            self.ui(self._set_busy, False)

        threading.Thread(target=work, daemon=True).start()

    def export_result(self):
        if not self.last_results:
            messagebox.showwarning("Chưa có kết quả", "Chưa có kết quả chấm điểm nào để lưu.")
            return
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
                                            initialfile="Evaluation_Results.xlsx")
        if not path:
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "Results"
        extra = []
        for res in self.last_results:
            extra += [k for k in res if k not in RESULT_COLS and k not in extra]
        cols = extra + RESULT_COLS  # cột riêng (StudentId, Group…) đứng đầu
        ws.append(cols)
        for res in self.last_results:
            ws.append([res.get(c, "") for c in cols])
        from openpyxl.utils import get_column_letter
        widths = {"Question": 50, "ReferenceMaterial": 50, "StudentAnswer": 50, "Feedback": 70,
                  "RawJSON": 40, "Error": 40, "BusinessId": 38}
        self._style_sheet(ws, {get_column_letter(i): widths.get(c, 14) for i, c in enumerate(cols, 1)})
        try:
            wb.save(path)
            messagebox.showinfo("Đã lưu kết quả", f"Đã lưu {len(self.last_results)} dòng tại:\n{path}")
        except OSError as e:
            messagebox.showerror("Không lưu được", f"{e}\n\nNếu file đang mở trong Excel, hãy đóng lại rồi thử lại.")
